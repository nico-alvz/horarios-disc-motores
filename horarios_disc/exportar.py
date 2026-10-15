"""Tablas Excel/CSV con conservación del libro original del cliente."""
import datetime as dt
import os

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .config import BLOQUES, DIA_CORTO, DIA_IDX, DIAS
from .etl import limpiar_secciones
from .tipos import Version
from .verificacion import detalle_preferencias


def rango_hora(p, h, dur): return f"{p.horas[BLOQUES[h]][0]}–{p.horas[BLOQUES[h + dur - 1]][1]}"

def horario_texto(g):
    """Notación del cliente: 'Ma F (V3-Rob), Lu DE (U8)'."""
    return ", ".join(f"{DIA_CORTO[DIAS[r.Dia]]} {''.join(BLOQUES[r.Hora:r.Hora + r.Dur])}" + (f" ({r.Sala})" if r.Sala else "")
                     for r in g.sort_values(["Dia", "Hora"]).itertuples())

SUPUESTOS = [
    "Factible bajo supuestos: NO equivale a 'validado para operación' hasta confirmar los puntos siguientes con el Encargado Docente.",
    "Disponibilidad de profesores hora: se infiere por los DÍAS en que hoy ya dictan (no hay disponibilidad por bloque en la planilla).",
    "Capacidad de salas sin dato del cliente: ESTIMADA con el mayor cupo que el cliente ya ubicó allí (ver hoja de salas del notebook).",
    "Ocupación externa de salas (ciclo básico, otros departamentos) y prioridad de ciencias básicas: no vienen en la planilla, no modeladas.",
    "Pertenencia de estudiantes: los paralelos (C2, T#, L##) se modelan como subgrupos del semestre; confirmar cohortes con el cliente.",
    "Banner: sin formato de importación acreditado; esta hoja es una guía de digitación. Campos como campus, método de instrucción o fechas no están acreditados.",
]

def tablas_version(p, v: Version, extra_res=()):
    etq, rank = v.etq, v.rank
    t = v.t; por_clase = {cid: g for cid, g in t.groupby("ClaseId")}
    est = p.secciones["Motivo"].replace("", "PROGRAMADA").values
    filas = []
    for _, c in p.clases.iterrows():
        g = por_clase[c.ClaseId].sort_values(["Dia", "Hora"])
        for i in c.Filas:
            s = p.secciones.loc[i]; fila_xl = i + 2
            code_ = s["Code"] if pd.notna(s["Code"]) else "SINCODIGO"
            estado = "Pendiente: crear NRC en Banner" if str(s["NRC"]).upper() == "ADD" else ("Pendiente: falta código" if pd.isna(s["Code"]) else "Lista")
            for k, r in enumerate(g.itertuples(), 1):
                filas.append({"Periodo": p.cfg.periodo, "ID sección": f"{p.cfg.periodo}|{code_}-{s['P']}|fila{fila_xl}", "ID sesión": f"{p.cfg.periodo}|{code_}-{s['P']}|fila{fila_xl}|S{k}",
                              "NRC": s["NRC"], "Código": s["Code"], "Materia": s["Materia"], "Sección": s["P"], "Asignatura": s["Asignatura"],
                              "Programa": s["Programa"], "Semestre": s["Semestre"], "Tipo": s["Tipo"], "Liga": s["Liga"], "Cupo": s["Cupo"],
                              "Día": DIAS[r.Dia], "Bloque": "".join(BLOQUES[r.Hora:r.Hora + r.Dur]),
                              "Hora inicio": p.horas[BLOQUES[r.Hora]][0], "Hora término": p.horas[BLOQUES[r.Hora + r.Dur - 1]][1],
                              "Sala": r.Sala or "", "Docente(s) y porcentajes (planilla)": s["Profesor"],
                              "RUT(s) (planilla)": s["RUT"], "Estado de preparación": estado, "Fila en planilla": fila_xl})
    guia = pd.DataFrame(filas)
    guia["_d"] = guia["Día"].map(DIA_IDX)
    guia = guia.sort_values(["Programa", "Semestre", "Asignatura", "Sección", "_d", "Hora inicio"], na_position="last").drop(columns="_d")
    ind = v.ind
    res = pd.DataFrame([("Versión", etq if v.cambios is not None else f"{etq} de {p.cfg.n_versiones}"), ("Periodo", p.cfg.periodo), ("Generado", dt.datetime.now().strftime("%d-%m-%Y %H:%M")),
                        ("Motor", "FET 7.10.5 (fet-cl)"), ("Semillas (6 componentes)", ", ".join(map(str, v.corrida["semillas"]))),
                        ("MODO_SALAS", p.cfg.modo_salas), ("DIAS_HORA_ACOTADOS", p.cfg.dias_hora_acotados), ("FIJAR_SABADO", p.cfg.fijar_sabado),
                        ("MAX_HUECOS_SEMANA", p.cfg.max_huecos_semana), ("MAX_VENTANAS_DOCENTE", p.cfg.max_ventanas_docente),
                        ("TIEMPO_LIMITE_S", p.cfg.tiempo_limite_s), ("Archivo de entrada", os.path.basename(p.xlsx))]
                       + list(ind.items()) + [("Ranking", rank)] + list(extra_res)
                       + [("Supuesto / pendiente", z) for z in SUPUESTOS], columns=["Parámetro / indicador", "Valor"])
    if v.restr:
        res = pd.concat([res, pd.DataFrame(
            [("Restricción aplicada", str(r)) for r in v.restr],
            columns=["Parámetro / indicador", "Valor"]
        )], ignore_index=True)
        detalle = detalle_preferencias(p, v.t, v.restr)
        res = pd.concat([res, pd.DataFrame(
            [("Cumplimiento de preferencia", f"{r['Restricción']}: {r['Cumplidas']} / {r['Sesiones afectadas']} sesiones ({r['%']} %)")
             for r in detalle.to_dict("records")],
            columns=["Parámetro / indicador", "Valor"]
        )], ignore_index=True)
    choq = pd.DataFrame({"Choque": v.viol or ["(ninguno: versión aprobable bajo los supuestos declarados)"]})
    return est, guia, res, choq

def estilo_hoja(ws, df, congelar=True):
    if congelar: ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B3C6F"); cell.alignment = Alignment(wrap_text=True, vertical="center")
    for j, col in enumerate(df.columns, 1):
        ancho = max([len(str(col))] + [len(str(z)) for z in df[col].head(300)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(9, ancho + 2), 60)

def df_a_hoja(wb, nombre, df):
    ws = wb.create_sheet(nombre)
    ws.append(list(df.columns))
    for fila in df.itertuples(index=False): ws.append([None if pd.isna(z) else (z.item() if hasattr(z, 'item') else z) for z in fila])   # pd.NA / NaN -> celda vacía
    estilo_hoja(ws, df)

COL_HORARIO = 25   # columna 'Horario' de la hoja 'asignaturas'

def exportar_xlsx(p, v: Version, ruta, extra_res=(), hojas_extra=None):
    """Copia del libro ORIGINAL (fórmulas y hojas intactas) con solo lo autorizado modificado."""
    est, guia, res, choq = tablas_version(p, v, extra_res)
    wb2 = openpyxl.load_workbook(p.xlsx)                 # data_only=False: conserva fórmulas
    ws = wb2["asignaturas"]
    assert ws.cell(1, COL_HORARIO).value == "Horario"
    n_col = ws.max_column
    ws.cell(1, n_col + 1, "Horario_original"); ws.cell(1, n_col + 2, "Estado_programacion")
    for r_ in range(2, ws.max_row + 1):
        ws.cell(r_, n_col + 1, ws.cell(r_, COL_HORARIO).value); ws.cell(r_, n_col + 2, est[r_ - 2])
    por_clase = {cid: g for cid, g in v.t.groupby("ClaseId")}
    for _, c in p.clases.iterrows():
        h = horario_texto(por_clase[c.ClaseId])
        for i in c.Filas:
            assert str(ws.cell(i + 2, 1).value) == str(p.raw.at[i, "NRC"]), f"fila desalineada: {i + 2}"
            ws.cell(i + 2, COL_HORARIO, h)
    for cell in ws[1][n_col:n_col + 2]: cell.font = Font(bold=True)
    df_a_hoja(wb2, "guia_digitacion", guia); df_a_hoja(wb2, "resumen", res); df_a_hoja(wb2, "choques", choq)
    if v.restr:
        restricciones = pd.DataFrame(
            [{"Restricción": str(r), "Objetivo": r.objetivo, "Referencia": r.ref,
              "Días": ", ".join(r.dias or []), "Bloques": "".join(r.bloques or []),
              "Dura": r.duro, "Peso": r.peso} for r in v.restr]
        )
        df_a_hoja(wb2, "restricciones", restricciones)
        df_a_hoja(wb2, "preferencias", detalle_preferencias(p, v.t, v.restr))
    extras = dict(hojas_extra or {})
    if v.cambios is not None:
        extras["cambios_vs_base"] = v.cambios
    for nombre_h, df_h in extras.items(): df_a_hoja(wb2, nombre_h, df_h)
    wb2.save(ruta)
    return guia

def verificar_integridad(p, ruta):
    """Compara celda a celda el archivo GUARDADO contra el original (fórmulas incluidas)."""
    a, b = openpyxl.load_workbook(p.xlsx), openpyxl.load_workbook(ruta)
    assert a.sheetnames == b.sheetnames[:len(a.sheetnames)], "cambió el orden o los nombres de las hojas originales"
    difs = 0
    for nombre in a.sheetnames:
        wa, wb_ = a[nombre], b[nombre]
        for fila in wa.iter_rows():
            for cel in fila:
                if nombre == "asignaturas" and cel.column == COL_HORARIO and cel.row > 1: continue   # única columna autorizada
                if cel.value != wb_[cel.coordinate].value: difs += 1
    assert difs == 0, f"{difs} celdas originales fueron alteradas"
    # reimportación: el libro exportado debe poder leerse otra vez con el mismo ETL
    re_sec = limpiar_secciones(pd.read_excel(ruta, sheet_name="asignaturas", dtype=object), p.academ)
    assert (re_sec["Motivo"] == "").sum() == (p.secciones["Motivo"] == "").sum(), "la reimportación cambia las secciones válidas"
    return True


def exportar_version(p, v: Version, carpeta) -> list[str]:
    """Exporta solo diagnósticos si la versión tiene choques; no imprime."""
    os.makedirs(carpeta, exist_ok=True)
    if not v.aprobable:
        ruta = os.path.join(carpeta, f"Choques_{v.etq}.csv")
        pd.DataFrame({"Choque": v.viol}).to_csv(ruta, index=False, encoding="utf-8-sig")
        return [ruta]
    prefijo = "Replan" if v.cambios is not None else "Horario"
    ruta = os.path.join(carpeta, f"{prefijo}_{v.etq}.xlsx")
    guia = exportar_xlsx(p, v, ruta)
    verificar_integridad(p, ruta)
    csv = os.path.join(carpeta, f"guia_digitacion_{v.etq}.csv")
    guia.to_csv(csv, index=False, encoding="utf-8-sig")
    v.xlsx, v.csv = ruta, csv
    return [ruta, csv]

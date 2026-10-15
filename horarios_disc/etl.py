"""ETL: planilla del Encargado Docente -> `Problema` (clases, actividades, docentes, salas)."""
import os
import re
import unicodedata
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

from .config import (ALIAS_SALA, BLOQUES, DIA_ABR, Config)

SESION_RE = re.compile(r"(Lu|Ma|Mi|Ju|Vi|Sa)\s+([A-G]{1,3})\s*(?:\(([^)]+)\))?")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]+", " ", s).split()


def clean_rut(v):
    if v is None or (isinstance(v, float) and pd.isna(v)): return None
    s = re.sub(r"[^0-9kK]", "", str(v)).upper()
    return s or None


def rut_valido(r):
    cuerpo, dv = r[:-1], r[-1]
    if not cuerpo.isdigit(): return False
    suma, m = 0, 2
    for d in reversed(cuerpo): suma += int(d) * m; m = 2 if m == 7 else m + 1
    esp = 11 - suma % 11
    return dv == ("0" if esp == 11 else "K" if esp == 10 else str(esp))


def canon_sala(r):
    r = r.strip(); return ALIAS_SALA.get(r.upper(), r)


def parse_horario(h):
    """'Ma F (V3), Mi E (V3)' -> [(dia, [bloques], sala)]"""
    if h is None or (isinstance(h, float) and pd.isna(h)): return []
    return [(DIA_ABR[d], list(bl), canon_sala(r) if r else None) for d, bl, r in SESION_RE.findall(str(h))]


def merge_adjacent(sesiones):
    """Bloques consecutivos, mismo día, sala y turno (A–C | D–G) -> una sola sesión doble."""
    out = []
    for dia, bl, sala in sesiones:
        bl = sorted(bl)
        if out and out[-1][0] == dia and out[-1][2] == sala and BLOQUES.index(bl[0]) == BLOQUES.index(out[-1][1][-1]) + 1 \
                and BLOQUES.index(bl[0]) != 3:
            out[-1] = (dia, out[-1][1] + bl, sala)
        else:
            out.append((dia, bl, sala))
    return out


def extraer_salas_y_bloques(wb):
    """Capacidades (títulos tipo 'Redes (Y1-005): 20') y horas de cada bloque, desde la hoja 'laboratorios'."""
    ws = wb["laboratorios"]; cap = {}
    for row in ws.iter_rows():
        for c in row:
            if c.column in (2, 12) and isinstance(c.value, str) and "Bloque" not in c.value:
                m = re.match(r"^\s*(.+?)\s*:\s*(\d+)", c.value)
                if not m: continue
                nombre, n = m.group(1), int(m.group(2))
                for tok in re.findall(r"Y\d-\d+", nombre): cap[canon_sala(tok)] = n
                if "Ladeco Rob" in nombre: cap["V2-Rob"] = n
                if "Ladeco Univ" in nombre: cap["V2-Univ"] = n
                if nombre.startswith("V3-Rob"): cap["V3-Rob"] = n
    horas = {}
    for r in ws.iter_rows(min_row=3, max_row=12, min_col=2, max_col=4, values_only=True):
        if r[0] in BLOQUES and hasattr(r[1], "strftime"): horas[r[0]] = (r[1].strftime("%H:%M"), r[2].strftime("%H:%M"))
    return cap, horas


def limpiar_secciones(raw, academ):
    a = raw.copy(); a.columns = [str(c) for c in a.columns]
    a = a.rename(columns={"Tipo.1": "Contrato"})
    for c in ["Code", "Asignatura"]:
        a[c] = a[c].astype("string").str.replace(r"\s+", " ", regex=True).str.strip()
    a["Asignatura"] = a["Asignatura"].fillna("(sin nombre)")
    a["Tipo"] = a["Tipo"].astype("string").str.strip()
    for c in ["Cupo", "Block", "Semestre"]: a[c] = pd.to_numeric(a[c], errors="coerce")
    a["Sesiones"] = a["Horario"].map(parse_horario)

    por_nombre = {tuple(norm(r["Nombre"])[:2]): clean_rut(r["RUT"]) for _, r in academ.iterrows()}
    IGNORAR = {"N. N."}   # docente por definir; las iniciales sueltas se descartan por len(n.split()) < 2

    def docentes(row):
        ruts = [clean_rut(x) for x in re.split(r"\s+[-+]\s+|\s*/\s*", str(row["RUT"]))] if pd.notna(row["RUT"]) else []
        nombres = re.split(r"\s*(?:\+|,)\s*", re.sub(r"\(\d+%\)", "", str(row["Profesor"]))) if pd.notna(row["Profesor"]) else []
        nombres = [n.strip() for n in nombres if n.strip()]
        out = []
        for i, n in enumerate(nombres):
            if len(n.split()) < 2 or n.upper() in IGNORAR: continue
            rut = ruts[i] if len(ruts) == len(nombres) else (ruts[0] if len(ruts) == 1 and len(nombres) == 1 else None)
            rut = rut or por_nombre.get(tuple(norm(n)[:2]))
            out.append((rut or "SINRUT-" + "".join(norm(n)), n))
        return out
    a["Docentes"] = a.apply(docentes, axis=1)

    a["Motivo"] = ""
    a.loc[a["Estado"].astype(str).str.upper().eq("INACTIVE"), "Motivo"] = "estado INACTIVE"
    a.loc[a["Resolución CAD"].astype(str).eq("No se dictará"), "Motivo"] = "CAD: no se dictará"
    a.loc[(a["Motivo"] == "") & (a["Sesiones"].map(len) == 0) & a["Block"].isna(), "Motivo"] = "sin horario ni bloques"
    return a


def construir_clases(sec):
    """Una *clase* = unidad que FET debe ubicar. Las secciones ligadas (mismo curso, docente y horario) se fusionan."""
    v = sec[sec["Motivo"] == ""].copy()
    v["Ses"] = v["Sesiones"].map(merge_adjacent)
    v["SesFallback"] = v.apply(lambda r: [(None, ["A"], None)] * int(r["Block"]) if not r["Ses"] and pd.notna(r["Block"]) else [], axis=1)

    def firma(r):
        return (r["Liga"] if pd.notna(r["Liga"]) else f"_{r.name}", r["P"], tuple(t[0] for t in r["Docentes"]),
                tuple((d, tuple(b)) for d, b, _ in r["Ses"]))
    v["_k"] = v.apply(firma, axis=1)
    filas = []
    for k, g in v.groupby("_k", sort=False):
        r0 = g.iloc[0]
        grupos = sorted({(p, int(sm)) for p, sm, t in zip(g["Programa"], g["Semestre"], g["Tipo"])
                         if pd.notna(sm) and pd.notna(p) and "/" not in str(p)
                         and t in ("Obligatorio", "Taller", "Proyecto", "Cátedra", "Laboratorio")})
        filas.append(dict(
            ClaseId=len(filas) + 1, Filas=list(g.index), NRC=" / ".join(map(str, g["NRC"])), Codigo=" / ".join(g["Code"].dropna().unique()),
            Asignatura=r0["Asignatura"], Seccion=r0["P"], Tipo=r0["Tipo"],
            Programas=" / ".join(sorted(set(g["Programa"].dropna()))), Cupo=int(g["Cupo"].fillna(0).sum()),
            Docentes=r0["Docentes"], Contrato=r0["Contrato"] if pd.notna(r0["Contrato"]) else "",
            Grupos=grupos, SesionesCliente=r0["Ses"] or r0["SesFallback"],
            SalasCliente=sorted({sa for ses in g["Ses"] for _, _, sa in ses if sa})))
    return pd.DataFrame(filas)


def construir_actividades(clases, fijar_sabado):
    acts, gid = [], 0
    for _, c in clases.iterrows():
        n = len(c.SesionesCliente); ids = []
        for k, (dia, bl, sala) in enumerate(c.SesionesCliente):
            gid += 1; ids.append(gid)
            acts.append(dict(Id=gid, ClaseId=c.ClaseId, Sesion=k + 1, Dur=len(bl), DiaCliente=dia,
                             BloqueCliente=BLOQUES.index(bl[0]), SalaCliente=sala,
                             Fijo=bool(fijar_sabado and dia == "Sábado"),
                             GrupoId=0 if n == 1 else ids[0], TotalDur=sum(len(s[1]) for s in c.SesionesCliente)))
    return pd.DataFrame(acts)


@dataclass
class Problema:
    """Todo lo que el ETL extrae de la planilla; es el único estado compartido del módulo."""
    cfg: Config
    xlsx: str
    raw: pd.DataFrame
    academ: pd.DataFrame
    wb: openpyxl.Workbook
    secciones: pd.DataFrame
    clases: pd.DataFrame
    acts: pd.DataFrame
    docentes: dict          # rut -> dict(rut, nombre, contrato, dias)
    salas: list
    cap_declarada: dict
    cap: dict               # capacidad efectiva (límite) por sala
    horas: dict             # bloque -> (inicio, fin)
    clase_por_id: dict = field(default_factory=dict)

    @property
    def horas_fet(self):
        return [f"{b} {self.horas[b][0]}-{self.horas[b][1]}" for b in BLOQUES]


def cargar_problema(xlsx, cfg=None):
    """Lee `datos cliente.xlsx` y devuelve el `Problema` listo para generar horarios."""
    cfg = cfg or Config()
    if not os.path.isfile(xlsx): raise FileNotFoundError(f"Falta la planilla: {xlsx}")
    raw = pd.read_excel(xlsx, sheet_name="asignaturas", dtype=object)
    academ = pd.read_excel(xlsx, sheet_name="academicos", dtype=object).dropna(how="all")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    cap_decl, horas = extraer_salas_y_bloques(wb)
    secciones = limpiar_secciones(raw, academ)
    clases = construir_clases(secciones)

    uso = {}
    for _, c in clases.iterrows():
        for sa in c.SalasCliente: uso[sa] = max(uso.get(sa, 0), c.Cupo)
    salas = sorted(set(uso) | set(cap_decl))
    cap = {s: cap_decl[s] if s in cap_decl else max(uso.get(s, 0), 10) for s in salas}   # sin dato: se ESTIMA

    docentes = {}
    for _, c in clases.iterrows():
        for rut, nombre in c.Docentes:
            d = docentes.setdefault(rut, dict(rut=rut, nombre=nombre, contrato=set(), dias=set()))
            if c.Contrato: d["contrato"].add(c.Contrato)
            d["dias"] |= {dia for dia, _, _ in c.SesionesCliente if dia}
    for d in docentes.values(): d["contrato"] = "/".join(sorted(d["contrato"])) or "s/i"

    p = Problema(cfg, xlsx, raw, academ, wb, secciones, clases, construir_actividades(clases, cfg.fijar_sabado),
                 docentes, salas, cap_decl, cap, horas)
    p.clase_por_id = {c.ClaseId: c for _, c in clases.iterrows()}
    p._uso = uso
    return p


def informe_etl(p):
    """Calidad del ETL: salas, RUT inválidos, docentes sin RUT y secciones excluidas (con su motivo)."""
    salas = pd.DataFrame({"Sala": p.salas, "Capacidad declarada": [p.cap_declarada.get(s) for s in p.salas],
                          "Mayor cupo ubicado por el cliente": [p._uso.get(s) for s in p.salas],
                          "Capacidad usada como límite": [p.cap[s] for s in p.salas]})
    salas["Origen"] = ["declarada" if s in p.cap_declarada else "ESTIMADA (confirmar)" for s in p.salas]
    salas["Alerta"] = [f"el cliente ya ubica cupo {p._uso[s]} > capacidad declarada {p.cap_declarada[s]}"
                       if s in p.cap_declarada and p._uso.get(s, 0) > p.cap_declarada[s] else "" for s in p.salas]
    sin_rut = [d["nombre"] for d in p.docentes.values() if d["rut"].startswith("SINRUT")]
    rut_malos = [(d["nombre"], d["rut"]) for d in p.docentes.values() if not d["rut"].startswith("SINRUT") and not rut_valido(d["rut"])]
    excl = p.secciones[p.secciones.Motivo != ""][["NRC", "Code", "Asignatura", "Motivo"]].reset_index(drop=True)
    return dict(salas=salas, sin_rut=sin_rut, rut_invalidos=rut_malos, excluidas=excl,
                resumen=f"{len(p.secciones)} secciones -> {(p.secciones.Motivo == '').sum()} válidas -> {len(p.clases)} clases "
                        f"({len(p.acts)} sesiones semanales) | {len(p.docentes)} docentes | {len(p.salas)} salas")
